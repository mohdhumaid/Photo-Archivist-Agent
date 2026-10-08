"""Step 4: Text — text layer first; OCR only when no text layer exists."""
from __future__ import annotations
import os
import shutil
import subprocess


def pdf_text(path: str) -> tuple[str, bool]:
    """Returns (text, has_layer)."""
    try:
        from pypdf import PdfReader
        r = PdfReader(path)
        texts = [(p.extract_text() or "") for p in r.pages]
        joined = "\n".join(texts).strip()
        return joined[:20000], bool(joined)
    except Exception:
        return "", False


def ooxml_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            from docx import Document
            d = Document(path)
            bits = ["\n".join(p.text for p in d.paragraphs)]
            # comments + headers/footers carry ref numbers & review notes
            try:
                bits += [c.text for p in d.paragraphs for c in getattr(p, "comments", []) or []]
            except Exception:
                pass
            try:
                for sect in d.sections:
                    for part in (sect.header.paragraphs + sect.footer.paragraphs):
                        if part.text.strip():
                            bits.append(part.text)
            except Exception:
                pass
            try:
                bits += [t.text for t in d.tables for r in t.rows for t in r.cells][:200]
            except Exception:
                pass
            return "\n".join(b for b in bits if b)[:20000]
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            out: list[str] = []
            for ws in wb.worksheets:
                out.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    out.append(" | ".join("" if v is None else str(v) for v in row))
                try:
                    out.append(f"definedNames: {','.join(str(n) for n in list(wb.defined_names)[:20])}")
                except Exception:
                    pass
            return "\n".join(out)[:20000]
        if ext == ".pptx":
            try:
                from pptx import Presentation
            except ImportError:
                return ""  # optional dep (pip install python-pptx)
            prs = Presentation(path)
            return "\n".join(sh.text for sl in prs.slides for sh in sl.shapes
                             if getattr(sh, "has_text_frame", False))[:20000]
    except Exception:
        return ""
    return ""


def ocr_image(path: str) -> str:
    """OCR for scanned docs / signage. Called only when no text layer exists.

    Uses the `tesseract` binary directly via subprocess.
    If tesseract fails or is missing, falls back to EasyOCR (if installed).
    Finally, falls back to the vision-LLM transcribe path.
    """
    # 1. Try Tesseract
    exe = shutil.which("tesseract")
    if exe:
        try:
            out = subprocess.run([exe, path, "stdout", "-l", "eng"],
                                 capture_output=True, text=True, timeout=120)
            text = (out.stdout or "")[:20000].strip()
            if text:
                return text
        except Exception:
            pass

    # 2. Try EasyOCR fallback
    try:
        import easyocr
        reader = easyocr.Reader(['en'], gpu=False, verbose=False)
        result = reader.readtext(path, detail=0)
        text = "\n".join(result).strip()
        if text:
            return text
    except Exception:
        pass

    # 3. Last resort: vision-LLM
    return _vision_ocr_hint(path)


def _vision_ocr_hint(path: str) -> str:
    """OCR last resort: prompt the vision-LLM to transcribe visible text.

    Runs only when vision_llm.enabled is true and the gateway is reachable;
    otherwise returns "" — OCR stays absent rather than invented.
    """
    try:
        import yaml
    except ImportError:
        return ""
    cfg = {}
    try:
        if os.path.exists("config.yaml"):
            with open("config.yaml") as f:
                cfg = yaml.safe_load(f) or {}
    except Exception:
        return ""
    from . import vllm as vllmmod
    if not vllmmod.enabled(cfg):
        return ""
    try:
        prime = {"file_name": os.path.basename(path), "file_type": "image",
                 "mime": "image/jpeg", "path_segments": [], "ocr_text": "",
                 "metadata": {}}
        out = vllmmod.describe(cfg, path, prime)
        if not isinstance(out, dict):
            return ""
        return str(out.get("visible_text") or out.get("caption") or "")[:4000]
    except Exception:
        return ""

